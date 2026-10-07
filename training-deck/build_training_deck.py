#!/usr/bin/env python3
"""Build the Zyda "Dashboard Training" deck for customer-service agents (Arabic + English, right-to-left).

    python build_training_deck.py                      # uses ./screenshots, writes ./output/...pptx
    python build_training_deck.py --list               # print every screenshot slot and what it should show
    python build_training_deck.py --placeholders-only  # blank template: every slot is a placeholder

The deck follows the order workflow: 1 login & brand selection, 2 customer data, 3 order processing,
4 checkout & payment, 5 order confirmation.  Text is Egyptian Arabic with the dashboard's English UI terms
kept as they appear on screen (written as **bold** in the specs below and drawn in the section colour).

How screenshots work
--------------------
Every screenshot slot has a file name (see --list), e.g. ``01_login``.  If a PNG/JPG with that name exists
in the screenshots folder it is placed on the slide together with its numbered call-outs; if not, a dashed
placeholder frame is drawn that says which file to drop in.  Re-run the script to refresh the deck.

The numbered circles on a screenshot match the numbered steps next to it.  Call-out positions are stored as
fractions of the *original* screenshot, so a re-captured screenshot of the same screen keeps its call-outs.

Requires: python-pptx (pip install python-pptx).  Pillow is installed together with it.
"""
from __future__ import annotations

import argparse
import math
import os
import re
import sys
from dataclasses import dataclass
from typing import Dict, List, Optional, Sequence, Tuple

from PIL import Image, ImageFont, features
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# --------------------------------------------------------------------------------------
# Design tokens
# --------------------------------------------------------------------------------------
SLIDE_W, SLIDE_H = 13.333, 7.5
MX = 0.8  # left/right margin
CONTENT_W = SLIDE_W - 2 * MX
CONTENT_TOP = 1.95
CONTENT_BOTTOM = 7.0
RTL = True  # Arabic-first layout: reading starts on the right

FONT = "Arial"  # Latin text
FONT_AR = "Arial"  # Arabic (complex-script) text - Arial ships Arabic glyphs on Windows and macOS
RUNNING_HEADER = "RING"  # small label centred at the top of every slide (as on the supplied cover)

INK = "141821"
MUTED = "5B6272"
FAINT = "8A90A0"
RULE = "DDE1EA"
BG = "F5F6FA"  # content slides
CARD = "FFFFFF"
BLACK = "000000"
WHITE = "FFFFFF"
WARN = "D92D20"

ZYDA_BLUE = "234DFB"  # sampled from the login screen
SECTIONS = {
    1: dict(ar="الدخول واختيار البراند", en="Login & Brand Selection", color=ZYDA_BLUE),
    2: dict(ar="بيانات العميل", en="Customer Data", color="0B8457"),
    3: dict(ar="تنفيذ الأوردر", en="Order Processing", color="E8590C"),
    4: dict(ar="المراجعة والدفع", en="Checkout & Payment", color="6D3EF5"),
    5: dict(ar="تأكيد الأوردر", en="Order Confirmation", color="D6336C"),
}

SMAX = 0.0125  # max inches per screenshot pixel (avoids blowing small crops up too far)

LATIN_FONT_FILES = {
    False: ["/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "C:/Windows/Fonts/arial.ttf",
            "/Library/Fonts/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf"],
    True: ["/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf", "C:/Windows/Fonts/arialbd.ttf",
           "/Library/Fonts/Arial Bold.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"],
}
ARABIC_FONT_FILES = {  # fonts with both Arabic and Latin glyphs, used only to measure mixed text
    False: ["C:/Windows/Fonts/arial.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"],
    True: ["C:/Windows/Fonts/arialbd.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
           "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"],
}

ARABIC = re.compile("[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]")


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def tint(hex_: str, f: float) -> str:
    """Mix a colour with white; f=0 -> colour, f=1 -> white."""
    r, g, b = (int(hex_[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(c + (255 - c) * f):02X}" for c in (r, g, b))


def shade(hex_: str, f: float) -> str:
    """Mix a colour with black."""
    r, g, b = (int(hex_[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(c * (1 - f)):02X}" for c in (r, g, b))


def plain(text: str) -> str:
    return text.replace("**", "")


# --------------------------------------------------------------------------------------
# Text measuring (so step lists can be laid out without guessing)
# --------------------------------------------------------------------------------------
_fonts: Dict[Tuple[bool, bool], Optional[ImageFont.FreeTypeFont]] = {}


def _measure_font(arabic: bool, bold: bool):
    key = (arabic, bold)
    if key not in _fonts:
        font = None
        files = ARABIC_FONT_FILES if arabic else LATIN_FONT_FILES
        engine = ImageFont.Layout.RAQM if (arabic and features.check("raqm")) else ImageFont.Layout.BASIC
        for path in files[bold]:
            if os.path.exists(path):
                try:
                    font = ImageFont.truetype(path, 200, layout_engine=engine)
                    break
                except OSError:
                    pass
        _fonts[key] = font
    return _fonts[key]


def text_width(text: str, size: float, bold: bool = False) -> float:
    """Width of one line of text in inches."""
    text = plain(text)
    arabic = bool(ARABIC.search(text))
    font = _measure_font(arabic, bold)
    if font is None:
        return len(text) * size * (0.56 if arabic else 0.54) / 72
    return font.getlength(text) / 200 * size / 72 * (1.04 if arabic else 1.0)


def count_lines(text: str, size: float, bold: bool, width: float) -> int:
    lines = 0
    for para in plain(text).split("\n"):
        cur, n = "", 1
        for w in para.split(" "):
            trial = (cur + " " + w).strip()
            if text_width(trial, size, bold) <= width or not cur:
                cur = trial
            else:
                n += 1
                cur = w
        lines += n
    return lines


# --------------------------------------------------------------------------------------
# Screenshot registry: key -> (file stem, (width px, height px), caption, what the slot should show)
# Sizes are the real screenshot sizes; they keep placeholder frames the same shape as the real thing.
# --------------------------------------------------------------------------------------
SHOT_TABLE = [
    ("login", "01_login", (1357, 601), "شاشة الـ Login",
     "صفحة Login to your store: خانة Email Address وخانة Password وزرار Login."),
    ("orders_incoming", "02_orders_incoming", (1100, 464), "الأوردر وهو بيوصل الداشبورد",
     "كارت أوردر جديد على الداشبورد: الصف اللي فوق (الوقت، رقم الأوردر، العميل، النوع، الدفع، الميعاد، الحالة) والأصناف والعميل."),
    ("order_full", "03_order_full", (1093, 557), "أوردر متسجّل كامل",
     "كارت أوردر كامل: الأصناف بالـ options، الحساب، العميل، العنوان، الخريطة، زرار Ready وسطر Order Accepted by."),
    ("order_customer_map", "04_order_customer_map", (528, 423), "بيانات العميل على الأوردر",
     "جزء العميل في الأوردر: الاسم، رقم الأوردر، الموبايل، العنوان، والخريطة بالفرع والعميل."),
    ("phone_entry", "05_phone_entry", (392, 153), "Select Customer",
     "شاشة Select Customer وخانة Phone Number (فيها +20)."),
    ("customer_new", "06_customer_new", (408, 290), "رقم مش متسجّل",
     "Select Customer بعد رقم مش متسجّل: This looks like a new customer وزرار Add Customer."),
    ("customer_details", "07_customer_details", (404, 348), "Complete Info (عميل جديد)",
     "فورم Complete Info: Phone Number و Full Name و Save."),
    ("customer_save", "08_customer_save", (392, 55), "Save",
     "زرار Save الأحمر بعد ما الاسم يبقى صح."),
    ("saved_addresses", "09_saved_addresses", (409, 636), "Select Address والعناوين المحفوظة",
     "الـ Checkout وشاشة Select Address مفتوحة: زرار Add new address والعناوين المحفوظة."),
    ("map_empty", "10_map_empty", (404, 643), "خريطة Order Mode",
     "Order Mode بتابات Delivery / Pickup وخريطة Enter Location فاضية."),
    ("map_search", "11_map_search", (441, 643), "دوّر واختار الاقتراح",
     "خريطة Order Mode بعد كتابة العنوان أو الإحداثيات وظهور الاقتراحات."),
    ("map_save", "12_map_save", (441, 643), "الـ Pin اتحط: Save",
     "خريطة Order Mode والـ pin متحط وزرار Save الأحمر."),
    ("not_deliverable", "13_not_deliverable", (392, 55), "برّه منطقة التوصيل",
     "زرار We don't deliver to this address."),
    ("search_results", "14_search_results", (416, 633), "نتايج البحث عن ice",
     "نتايج البحث في المنيو: الاسم والسعر والوصف والصورة لكل صنف."),
    ("menu_search", "15_menu_search", (402, 122), "Search و Menu",
     "خانة Search فوق المنيو."),
    ("basket_checkout", "16_basket_checkout", (397, 572), "الباسكت وزرار Checkout",
     "الباسكت (Your Items) وزراير Add Items و Checkout تحت."),
    ("checkout_bar", "17_checkout_bar", (380, 96), "الشريط اللي تحت بعد ما تضيف صنف",
     "الشريط اللي تحت: خط الخصم و Add Items و Checkout."),
    ("checkout_add_address", "18_checkout_add_address", (391, 218), "الـ Checkout: Please add an address",
     "جزء Deliver to في الـ Checkout ومكتوب Please add an address بالأحمر."),
    ("add_new_address_btn", "19_add_new_address_btn", (380, 55), "Add new address",
     "زرار Add new address الأحمر."),
    ("map_second_location", "20_map_second_location", (409, 643), "اللوكيشن تاني مرة",
     "شاشة اللوكيشن اللي بتفتح من Add new address: Enter Location والخريطة و Save."),
    ("edit_location", "21_edit_location", (380, 128), "Edit Location",
     "شريط Edit Location فوق فورم العنوان."),
    ("address_form", "22_address_form", (407, 643), "فورم العنوان",
     "فورم العنوان: House / Apartment / Office و City و Area و Street و Building و Floor و Apartment Number و Notes و Add new address."),
    ("item_options", "23_item_options", (407, 643), "شاشة الصنف",
     "صنف مفتوح: الاسم والوصف و Size وجروب Required والكمية وزرار Add Item."),
    ("basket_items", "24_basket_items", (354, 132), "الصنف في Your Items",
     "الصنف جوه الباسكت بالاختيار بتاعه وزراير الكمية."),
    ("order_item_options", "25_order_item_options", (527, 552), "الـ Extras على الأوردر",
     "أصناف أوردر فيها جروبات options و Extra بسعره."),
    ("sold_out_list", "26_sold_out_list", (407, 643), "كاتيجوري فيها أصناف Sold out",
     "كاتيجوري فيها صنف عادي وأصناف Sold out رمادي وصنف Schedule بس."),
    ("sold_out_item", "27_sold_out_item", (392, 105), "Sold out",
     "صنف واحد عليه Sold out."),
    ("pickup_branches", "28_pickup_branches", (409, 643), "Pickup: Select Branch",
     "Order Mode و Pickup مختار وليستة Select Branch (الاسم والمسافة والعنوان)."),
    ("pickup_branch_card", "29_pickup_branch_card", (367, 164), "الفرع المختار",
     "كارت Pickup Branch بالـ pin على الخريطة."),
    ("pickup_time", "30_pickup_time", (392, 157), "ميعاد الاستلام",
     "Pickup مختار: Picking up from + الفرع و Ready in 15 minutes - ASAP وزرار Schedule slot."),
    ("payment_place_order", "31_payment_place_order", (367, 163), "Pay with: Cash",
     "جزء Pay with (Cash / Credit Card) وزرار Place Order."),
    ("payment_options", "32_payment_options", (361, 110), "Pay with",
     "جزء Pay with (Cash / Credit Card)."),
    ("checkout_full", "33_checkout_full", (400, 638), "صفحة الـ Checkout",
     "صفحة الـ Checkout كلها: Deliver to والوقت و Save on this order (Voucher) و Order Summary و Place Order."),
    ("voucher_box", "34_voucher_box", (367, 76), "خانة الـ Voucher",
     "خانة Save on this order: Voucher و Apply."),
    ("voucher_code", "35_voucher_code", (364, 68), "كود الـ Voucher",
     "خانة Have a code? للـ Voucher."),
    ("summary_delivery", "36_summary_delivery", (367, 143), "Order Summary · Delivery",
     "Order Summary لأوردر Delivery: Subtotal و Delivery Services و Service Fee و Total."),
    ("summary_pickup", "37_summary_pickup", (367, 120), "Order Summary · Pickup",
     "Order Summary لأوردر Pickup (من غير Delivery Services)."),
    ("customer_existing", "38_customer_existing", (408, 290), "عميل مسجّل",
     "Select Customer بعد رقم متسجّل: بروفايل العميل بيظهر وتقدر تكمّل."),
    ("schedule_slot_picker", "40_schedule_slot_picker", (392, 290), "Schedule slot",
     "شاشة Schedule slot اللي بتختار منها يوم وساعة الاستلام."),
    ("pickup_order_recorded", "41_pickup_order_recorded", (520, 250), "أوردر Pickup على الداشبورد",
     "كارت أوردر Pickup على الداشبورد: الفرع ويوم وساعة الاستلام."),
    ("cashback_checkout", "42_cashback_checkout", (352, 316), "Credit Card: Cashback و Send Cart Link",
     "الـ Checkout و Credit Card مختار: سطر Cashback بالمبلغ وزرار Send Cart Link."),
    ("special_instructions", "43_special_instructions", (366, 566), "خانة Special Instructions",
     "صنف مفتوح وتحته خانة SPECIAL INSTRUCTIONS."),
    ("search_first_letters", "44_search_first_letters", (352, 416), "البحث بأول حروف: RAN",
     "البحث بأول كام حرف (مثلاً RAN) والنتايج اللي بتظهر."),
    ("switch_store", "45_switch_store", (1362, 630), "Switch Store",
     "صفحة Switch Store: You're now logged in to + البراند، Search using store name، وليستة البراندات."),
    ("store_switcher", "46_store_switcher", (296, 120), "اسم البراند في القائمة",
     "اسم البراند الحالي فوق في القائمة الجانبية."),
    ("last_order", "47_last_order", (409, 300), "آخر أوردر فوق الـ Categories",
     "آخر أوردر للعميل وهو ظاهر فوق الـ Categories في المنيو."),
]
SHOTS = {k: dict(file=f, size=s, caption=c, show=d) for k, f, s, c, d in SHOT_TABLE}


@dataclass
class Use:
    """One placement of a screenshot slot on a slide."""

    key: str
    marks: Sequence[Tuple[int, float, float]] = ()  # (step number, x, y) as fractions of the ORIGINAL image
    boxes: Sequence[Tuple[float, float, float, float]] = ()  # highlight rectangles, original-image fractions
    crop: Optional[Tuple[float, float, float, float]] = None  # visible part (x0, y0, x1, y1), original fractions
    cap: Optional[str] = None  # caption override ('' hides the caption)


def U(key, marks=(), boxes=(), crop=None, cap=None):
    return Use(key, tuple(marks), tuple(boxes), crop, cap)


# --------------------------------------------------------------------------------------
# Build context
# --------------------------------------------------------------------------------------
class Ctx:
    def __init__(self, shots_dir: str, placeholders_only: bool, header: str, logo: Optional[str]):
        self.shots_dir = shots_dir
        self.placeholders_only = placeholders_only
        self.header = header
        self.logo = logo
        self.found: List[str] = []
        self.missing: List[str] = []
        self._dims: Dict[str, Tuple[int, int]] = {}
        self.section_first_page: Dict[int, int] = {}
        self.color = ZYDA_BLUE  # colour of the section being drawn

    def path(self, key: str) -> Optional[str]:
        if self.placeholders_only:
            return None
        stem = SHOTS[key]["file"]
        for ext in (".png", ".jpg", ".jpeg", ".PNG", ".JPG", ".JPEG"):
            p = os.path.join(self.shots_dir, stem + ext)
            if os.path.exists(p):
                return p
        return None

    def full_size(self, key: str) -> Tuple[int, int]:
        p = self.path(key)
        if p:
            if p not in self._dims:
                with Image.open(p) as im:
                    self._dims[p] = im.size
            return self._dims[p]
        return SHOTS[key]["size"]

    def dims(self, use: Use) -> Tuple[float, float]:
        w, h = self.full_size(use.key)
        if use.crop:
            x0, y0, x1, y1 = use.crop
            return w * (x1 - x0), h * (y1 - y0)
        return float(w), float(h)

    def note(self, key: str, found: bool) -> None:
        bucket = self.found if found else self.missing
        if key not in bucket:
            bucket.append(key)


# --------------------------------------------------------------------------------------
# Drawing helpers.  Layout code works in *logical* coordinates (x measured from the reading-start
# edge); with RTL on, every helper mirrors x unless raw=True.
# --------------------------------------------------------------------------------------
def I(v: float) -> Emu:
    return Inches(v)


def mx(x: float, w: float) -> float:
    return SLIDE_W - x - w if RTL else x


def _align(a: str, raw: bool) -> PP_ALIGN:
    if RTL and not raw:
        a = {"l": "r", "r": "l"}.get(a, a)
    return {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[a]


def _runs(text: str):
    """Split '**bold**' markup into (text, is_highlight) runs."""
    parts = text.split("**")
    return [(t, i % 2 == 1) for i, t in enumerate(parts) if t]


def _style_run(r, text, size, bold, italic, color):
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = FONT
    rpr = r._r.get_or_add_rPr()
    if ARABIC.search(text):
        rpr.set("lang", "ar-EG")
    latin = rpr.find(qn("a:latin"))
    cs = rpr.find(qn("a:cs"))
    if cs is None:
        cs = rpr.makeelement(qn("a:cs"), {})
        latin.addnext(cs)
    cs.set("typeface", FONT_AR)


def add_text(slide, x, y, w, h, paras, anchor="t", wrap=True, raw=False):
    """paras: list of dicts made with P()."""
    box = slide.shapes.add_textbox(I(x if raw else mx(x, w)), I(y), I(w), I(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = _align(p.get("align", "l"), raw)
        if ARABIC.search(p["text"]):
            para._p.get_or_add_pPr().set("rtl", "1")
        if p.get("after") is not None:
            para.space_after = Pt(p["after"])
        para.line_spacing = p.get("spacing") or (1.1 if ARABIC.search(p["text"]) else 1.0)
        for text, hl in _runs(p["text"]):
            r = para.add_run()
            _style_run(r, text, p["size"], p["bold"] or hl, p.get("italic", False),
                       (p.get("hl") or p["color"]) if hl else p["color"])
    return box


def P(text, size=12, bold=False, color=INK, align="l", after=None, spacing=None, italic=False, hl=None):
    return dict(text=text, size=size, bold=bold, color=color, align=align, after=after, spacing=spacing,
                italic=italic, hl=hl)


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, rounded=0.0, dash=False, shape=None, raw=False):
    st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(st, I(x if raw else mx(x, w)), I(y), I(w), I(h))
    if rounded and shape is None:
        s.adjustments[0] = min(0.5, rounded / max(0.01, min(w, h)))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def add_line(slide, x1, y1, x2, y2, color=RULE, w=0.75):
    if RTL:
        x1, x2 = SLIDE_W - x1, SLIDE_W - x2
    ln = slide.shapes.add_connector(1, I(x1), I(y1), I(x2), I(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(w)
    return ln


def badge(slide, cx, cy, n, d=0.3, size=11, fill=None, ring=True, raw=False):
    if not raw and RTL:
        cx = SLIDE_W - cx
    o = add_rect(slide, cx - d / 2, cy - d / 2, d, d, fill=fill or ZYDA_BLUE, line=WHITE if ring else None,
                 line_w=1.25, shape=MSO_SHAPE.OVAL, raw=True)
    tf = o.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _style_run(r, str(n), size, True, False, WHITE)
    return o


ARROW = "←" if RTL else "→"
CHEVRON = "‹" if RTL else "›"


# --------------------------------------------------------------------------------------
# Slide chrome
# --------------------------------------------------------------------------------------
def new_slide(prs, ctx: Ctx, page: int, notes: str = "", bg: str = BG, header_color: str = FAINT, bar: Optional[str] = None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    if bar:
        add_rect(s, 0, 0, SLIDE_W, 0.09, fill=bar, raw=True)
    if ctx.header:
        add_text(s, SLIDE_W / 2 - 1.5, 0.3, 3.0, 0.2, [P(ctx.header, 8, color=header_color, align="c")], raw=True)
    add_text(s, SLIDE_W - MX - 0.8, 0.3, 0.8, 0.2, [P(str(page), 8, color=header_color, align="r")])
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def content_slide(prs, ctx, page, spec):
    sec = spec.get("sec")
    ctx.color = SECTIONS[sec]["color"] if sec else ZYDA_BLUE
    s = new_slide(prs, ctx, page, spec.get("notes", ""), bar=ctx.color)
    title_block(s, ctx, spec["title"], spec.get("en", ""), sec)
    return s


def title_block(slide, ctx, title: str, en: str, sec: Optional[int]):
    tw = 8.7
    lines = count_lines(title, 26, True, tw * 1.12)  # the measuring font runs wider than Arial
    add_text(slide, MX, 0.62, tw, 0.5 * lines + 0.1, [P(title, 26, bold=True, hl=ctx.color)])
    if en:
        add_text(slide, MX, 0.62 + 0.47 * lines + 0.06, tw, 0.3, [P(en.upper(), 10.5, bold=True, color=ctx.color)])
    if sec:
        x = 10.1
        add_text(slide, x, 0.66, SLIDE_W - MX - x, 0.7, [
            P(f"SECTION {sec:02d}", 8, bold=True, color=ctx.color, after=1),
            P(SECTIONS[sec]["ar"], 11, bold=True, color=INK, after=0),
            P(SECTIONS[sec]["en"], 9, color=MUTED),
        ])


def wordmark(slide, cx, cy, em_pt: float, ctx: Ctx):
    """Zyda wordmark ◂ zyda ▸ drawn as vector shapes - or a supplied logo image.  Not mirrored."""
    if ctx.logo and os.path.exists(ctx.logo):
        with Image.open(ctx.logo) as im:
            ar = im.width / im.height
        h = em_pt / 72 * 1.1
        slide.shapes.add_picture(ctx.logo, I(cx - h * ar / 2), I(cy - h / 2), I(h * ar), I(h))
        return
    em = em_pt / 72
    tw = text_width("zyda", em_pt, True)
    tri_h, gap = em * 0.50, em * 0.17
    tri_w = tri_h * 0.86
    total = tri_w + gap + tw + gap + tri_w
    x0 = cx - total / 2
    top = cy - em * 0.645  # puts the x-height centre of the text on cy
    add_text(slide, x0 + tri_w + gap - 0.02, top, tw + 0.3, em * 1.3, [P("zyda", em_pt, bold=True, color=WHITE, spacing=1.0)], wrap=False, raw=True)
    # an isosceles triangle points up; rotating 270 / 90 degrees makes it point left / right.
    for rot, slot_cx in ((270, x0 + tri_w / 2), (90, x0 + total - tri_w / 2)):
        t = add_rect(slide, slot_cx - tri_h / 2, cy - tri_w / 2, tri_h, tri_w, fill=WHITE, shape=MSO_SHAPE.ISOSCELES_TRIANGLE, raw=True)
        t.rotation = rot
    add_text(slide, x0 + total + 0.04, cy - em * 0.62, 0.4, 0.2, [P("TM", max(6, em_pt * 0.1), color=WHITE)], wrap=False, raw=True)


# --------------------------------------------------------------------------------------
# Screenshots: layout + drawing
# --------------------------------------------------------------------------------------
def fit_size(text: str, width: float, start=14.0, floor=9.0) -> float:
    size = start
    while size > floor and text_width(text, size, True) > width:
        size -= 0.5
    return size


def draw_placeholder(slide, ctx, key: str, rx, y, w, h):
    shot = SHOTS[key]
    add_rect(slide, rx, y, w, h, fill=tint(ctx.color, 0.93), line=ctx.color, line_w=1.25, dash=True, rounded=0.06, raw=True)
    fname = shot["file"] + ".png"
    if h < 1.0 or w < 2.2:
        add_text(slide, rx + 0.1, y, w - 0.2, h, [P(f"PLACEHOLDER  ·  {fname}", 9, bold=True, color=ctx.color, align="c")], anchor="m", raw=True)
        return
    add_text(slide, rx + 0.25, y + 0.1, w - 0.5, h - 0.2, [
        P("مكان الصورة · SCREENSHOT PLACEHOLDER", 9, bold=True, color=ctx.color, align="c", after=4),
        P(fname, fit_size(fname, w - 0.5), bold=True, color=INK, align="c", after=6),
        P(shot["show"], 10.5, color=MUTED, align="c"),
    ], anchor="m", raw=True)


def draw_use(slide, ctx: Ctx, use: Use, x, y, w, h, caption_h: float = 0.0):
    rx = mx(x, w)  # real position; everything inside the picture is drawn raw (pictures are never mirrored)
    p = ctx.path(use.key)
    ctx.note(use.key, bool(p))
    if p:
        pic = slide.shapes.add_picture(p, I(rx), I(y), I(w), I(h))
        if use.crop:
            x0, y0, x1, y1 = use.crop
            pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = x0, y0, 1 - x1, 1 - y1
        pic.line.color.rgb = rgb("C9CED9")
        pic.line.width = Pt(0.75)

        def local(fx, fy):
            if use.crop:
                x0, y0, x1, y1 = use.crop
                return (fx - x0) / (x1 - x0), (fy - y0) / (y1 - y0)
            return fx, fy

        for bx0, by0, bx1, by1 in use.boxes:
            lx0, ly0 = local(bx0, by0)
            lx1, ly1 = local(bx1, by1)
            lx0, ly0, lx1, ly1 = max(0, lx0), max(0, ly0), min(1, lx1), min(1, ly1)
            if lx1 > lx0 and ly1 > ly0:
                add_rect(slide, rx + lx0 * w, y + ly0 * h, (lx1 - lx0) * w, (ly1 - ly0) * h,
                         line=ctx.color, line_w=2.25, rounded=0.05, raw=True)
        for n, fx, fy in use.marks:
            lx, ly = local(fx, fy)
            if -0.02 <= lx <= 1.02 and -0.02 <= ly <= 1.02:
                badge(slide, rx + lx * w, y + ly * h, n, d=0.27, size=10.5, fill=ctx.color, raw=True)
    else:
        draw_placeholder(slide, ctx, use.key, rx, y, w, h)
    cap = SHOTS[use.key]["caption"] if use.cap is None else use.cap
    if cap and caption_h:
        cw = max(w, 1.8)
        add_text(slide, rx + w - cw if RTL else rx, y + h + 0.06, cw, caption_h,
                 [P(cap, 9, color=MUTED, italic=not ARABIC.search(cap), align="r" if RTL else "l")], raw=True)


def place(slide, ctx: Ctx, rows: List[List[Use]], box, flow: str = "", gap_h=0.3, gap_v=0.32, cap_h=0.3):
    """Place rows of screenshots inside box=(x, y, w, h) (logical coordinates).  Each row is scaled on its own,
    then everything shrinks together if the stack is too tall.  Rows read from the start edge."""
    bx, by, bw, bh = box
    info = []
    for row in rows:
        dims = [ctx.dims(u) for u in row]
        sumw = sum(d[0] for d in dims)
        maxh = max(d[1] for d in dims)
        s = min((bw - (len(row) - 1) * gap_h) / sumw, SMAX)
        has_cap = any((SHOTS[u.key]["caption"] if u.cap is None else u.cap) for u in row)
        info.append(dict(row=row, dims=dims, s=s, maxh=maxh, cap=cap_h if has_cap else 0.0))
    fixed = (len(rows) - 1) * gap_v + sum(r["cap"] for r in info)
    tot = sum(r["s"] * r["maxh"] for r in info)
    f = min(1.0, max(0.2, (bh - fixed) / tot)) if tot else 1.0
    y = by + max(0.0, (bh - tot * f - fixed) * 0.4)
    for ri, r in enumerate(info):
        s = r["s"] * f
        row_w = sum(d[0] for d in r["dims"]) * s + (len(r["row"]) - 1) * gap_h
        x = bx + (bw - row_w) / 2
        row_h = r["maxh"] * s
        for ui, (use, d) in enumerate(zip(r["row"], r["dims"])):
            w, h = d[0] * s, d[1] * s
            yy = y + (row_h - h) / 2
            draw_use(slide, ctx, use, x, yy, w, h, r["cap"])
            if "h" in flow and ui < len(r["row"]) - 1:
                add_text(slide, x + w, yy + h / 2 - 0.16, gap_h, 0.3, [P(ARROW, 16, color=ctx.color, align="c")], anchor="m")
            x += w + gap_h
        if "v" in flow and ri < len(info) - 1:
            add_text(slide, bx, y + row_h + r["cap"] - 0.02, bw, gap_v, [P("↓", 15, color=ctx.color, align="c")], anchor="m")
        y += row_h + r["cap"] + gap_v


# --------------------------------------------------------------------------------------
# Banners and step lists
# --------------------------------------------------------------------------------------
def banner_style(kind: str, color: str):
    return {
        "tip": ("نصيحة · TIP", color),
        "warn": ("خد بالك · WATCH OUT", WARN),
        "must": ("دايماً · ALWAYS", WARN),
        "note": ("معلومة · GOOD TO KNOW", MUTED),
    }[kind]


def banner_height(text: str, width: float, size=12) -> float:
    return count_lines(text, size, False, width - 0.45) * size * 1.35 / 72 + 0.55


def draw_banner(slide, ctx, kind: str, text: str, x, y, w, size=12):
    label, col = banner_style(kind, ctx.color)
    h = banner_height(text, w, size)
    add_rect(slide, x, y, w, h, fill=tint(col, 0.9))
    add_rect(slide, x, y, 0.07, h, fill=col)
    add_text(slide, x + 0.25, y + 0.11, w - 0.4, 0.22, [P(label, 8.5, bold=True, color=col)])
    add_text(slide, x + 0.25, y + 0.34, w - 0.4, h - 0.36, [P(text, size, hl=col)])
    return h


STEP_SIZES = ((15, 12.5), (14, 11.5), (13, 11), (12, 10))
STEP_GAP = 0.17
BADGE = 0.32


def draw_steps(slide, ctx, steps, x, y, width, avail_h):
    tw = width - BADGE - 0.15
    chosen = None
    for t_sz, d_sz in STEP_SIZES:
        hs = []
        for title, detail in steps:
            h = count_lines(title, t_sz, True, tw) * t_sz * 1.3 / 72
            if detail:
                h += 0.04 + count_lines(detail, d_sz, False, tw) * d_sz * 1.32 / 72
            hs.append(h)
        chosen = (t_sz, d_sz, hs)
        if sum(hs) + STEP_GAP * (len(steps) - 1) <= avail_h:
            break
    t_sz, d_sz, hs = chosen
    cy = y
    for i, ((title, detail), h) in enumerate(zip(steps, hs), 1):
        badge(slide, x + BADGE / 2, cy + t_sz * 1.3 / 72 / 2 + 0.01, i, d=BADGE, size=11.5, fill=ctx.color, ring=False)
        paras = [P(title, t_sz, bold=True, hl=ctx.color, after=2 if detail else 0)]
        if detail:
            paras.append(P(detail, d_sz, color=MUTED, hl=INK))
        add_text(slide, x + BADGE + 0.15, cy - 0.02, tw, h + 0.1, paras)
        cy += h + STEP_GAP


def left_column(slide, ctx, spec, col_w):
    """Steps, optional inset screenshots and a banner stacked in the reading-start column."""
    x, top, bottom = MX, CONTENT_TOP, CONTENT_BOTTOM
    banner = spec.get("banner")
    if banner:
        bh = banner_height(banner[1], col_w)
        draw_banner(slide, ctx, banner[0], banner[1], x, bottom - bh, col_w)
        bottom -= bh + 0.2
    plans, inset_h = [], 0.0
    for u in spec.get("inset", []):
        w_px, h_px = ctx.dims(u)
        s_ = min(col_w / w_px, 0.0105)
        cap = SHOTS[u.key]["caption"] if u.cap is None else u.cap
        ch = 0.3 if cap else 0.0
        plans.append((u, w_px * s_, h_px * s_, ch))
        inset_h += h_px * s_ + ch + 0.12
    draw_steps(slide, ctx, spec["steps"], x, top, col_w, bottom - inset_h - (0.1 if plans else 0) - top)
    iy = bottom - inset_h + 0.12
    for u, w, h, ch in plans:
        draw_use(slide, ctx, u, x, iy, w, h, ch)  # logical x = start edge of the column
        iy += h + ch + 0.12


# --------------------------------------------------------------------------------------
# Slide kinds
# --------------------------------------------------------------------------------------
def k_cover(prs, ctx, page, spec):
    s = new_slide(prs, ctx, page, spec.get("notes", ""), bg=BLACK, header_color=WHITE)
    add_text(s, 0.85, 1.38, 6.5, 1.2, [P("DASHBOARD", 28, bold=True, color=WHITE, spacing=1.0),
                                       P("TRAINING", 28, bold=True, color=WHITE, spacing=1.0)], raw=True)
    add_text(s, 0.85, 2.38, 6.5, 0.5, [P("تدريب الداشبورد للـ Agents", 18, bold=True, color=tint(ZYDA_BLUE, 0.35), align="l")], raw=True)
    add_text(s, 10.06, 1.52, 2.6, 0.5, [P("Step by step", 10, color=WHITE), P("خطوة بخطوة", 10, color=WHITE, align="l")], raw=True)
    wordmark(s, SLIDE_W / 2, 4.15, 100, ctx)
    w = SLIDE_W / len(SECTIONS)
    for i, sec in enumerate(SECTIONS.values()):
        add_rect(s, i * w, SLIDE_H - 0.14, w, 0.14, fill=sec["color"], raw=True)
    return s


def k_agenda(prs, ctx, page, spec):
    ctx.color = ZYDA_BLUE
    s = new_slide(prs, ctx, page, spec.get("notes", ""), bar=ZYDA_BLUE)
    title_block(s, ctx, spec["title"], spec["en"], None)
    y, row_h = 1.95, 0.92
    for i, (sec, blurb) in enumerate(spec["items"]):
        col = SECTIONS[sec]["color"]
        yy = y + i * row_h
        add_rect(s, MX, yy, CONTENT_W, row_h - 0.14, fill=CARD, line=RULE, line_w=0.75, rounded=0.08)
        add_rect(s, MX, yy, 0.95, row_h - 0.14, fill=col, rounded=0.08)
        add_text(s, MX, yy, 0.95, row_h - 0.14, [P(f"{sec:02d}", 22, bold=True, color=WHITE, align="c")], anchor="m")
        add_text(s, MX + 1.2, yy + 0.12, 4.0, 0.4, [P(SECTIONS[sec]["ar"], 17, bold=True)])
        add_text(s, MX + 1.2, yy + 0.47, 4.0, 0.3, [P(SECTIONS[sec]["en"], 10, color=col, bold=True)])
        add_text(s, MX + 5.4, yy + 0.12, 5.2, 0.56, [P(blurb, 12, color=MUTED, hl=INK)], anchor="m")
        add_text(s, CONTENT_W + MX - 1.1, yy + 0.12, 0.9, 0.56, [P(f"سلايد {ctx.section_first_page[sec]}", 11, color=FAINT, align="r")], anchor="m")
    return s


def k_divider(prs, ctx, page, spec):
    sec = spec["sec"]
    col = SECTIONS[sec]["color"]
    s = new_slide(prs, ctx, page, spec.get("notes", ""), bg=col, header_color=WHITE)
    add_rect(s, SLIDE_W - 4.6 if not RTL else 0, 0, 4.6, SLIDE_H, fill=shade(col, 0.12), raw=True)
    add_text(s, MX, 1.6, 4, 1.6, [P(f"{sec:02d}", 96, bold=True, color=WHITE, spacing=0.9)])
    add_text(s, MX, 3.25, 7.6, 1.0, [P(SECTIONS[sec]["ar"], 40, bold=True, color=WHITE)])
    add_text(s, MX, 4.2, 7.6, 0.5, [P(SECTIONS[sec]["en"].upper(), 16, bold=True, color=tint(col, 0.6))])
    y = 5.15
    for item in spec["items"]:
        w = text_width(item, 13, True) + 0.5
        add_rect(s, MX, y, w, 0.42, fill=shade(col, 0.25), rounded=0.2)
        add_text(s, MX, y, w, 0.42, [P(item, 13, bold=True, color=WHITE, align="c")], anchor="m")
        y += 0.52
    return s


def k_statement(prs, ctx, page, spec):
    s = new_slide(prs, ctx, page, spec.get("notes", ""), bg=WARN, header_color=WHITE)
    sec = spec.get("sec")
    if sec:
        add_text(s, MX, 0.7, 7, 0.4, [P(f"SECTION {sec:02d}  ·  {SECTIONS[sec]['ar']}", 10, bold=True, color=tint(WARN, 0.6))])
    add_text(s, MX, 1.7, 11.7, 1.3, [P(spec["headline"], 54, bold=True, color=WHITE)])
    add_text(s, MX, 2.95, 11.7, 0.5, [P(spec["en"].upper(), 20, bold=True, color=tint(WARN, 0.55))])
    add_text(s, MX, 3.6, 11.7, 0.5, [P(spec["sub"], 20, color=WHITE)])
    add_text(s, MX, 4.6, 6, 0.3, [P("ده يشمل · THIS INCLUDES", 10, bold=True, color=tint(WARN, 0.6))])
    n = len(spec["chips"])
    cw = (CONTENT_W - (n - 1) * 0.2) / n
    for i, chip in enumerate(spec["chips"]):
        x = MX + i * (cw + 0.2)
        add_rect(s, x, 5.0, cw, 0.95, fill=WHITE, rounded=0.1)
        add_text(s, x + 0.2, 5.0, 0.45, 0.95, [P("×", 28, bold=True, color=WARN, align="c")], anchor="m")
        add_text(s, x + 0.7, 5.0, cw - 0.85, 0.95, [P(chip[0], 16, bold=True, after=0), P(chip[1], 10, color=MUTED)], anchor="m")
    return s


def k_journey(prs, ctx, page, spec):
    ctx.color = ZYDA_BLUE
    s = new_slide(prs, ctx, page, spec.get("notes", ""), bar=ZYDA_BLUE)
    title_block(s, ctx, spec["title"], spec["en"], None)
    steps = spec["steps"]
    n = len(steps)
    gap = 0.22
    cw = (CONTENT_W - (n - 1) * gap) / n
    y0, h = 2.05, 3.6
    for i, (sec, title, en, desc) in enumerate(steps):
        col = SECTIONS[sec]["color"]
        x = MX + i * (cw + gap)
        add_rect(s, x, y0, cw, h, fill=col, rounded=0.12)
        add_text(s, x + 0.22, y0 + 0.2, 1, 0.7, [P(str(i + 1), 34, bold=True, color=WHITE)])
        add_text(s, x + 0.22, y0 + 1.0, cw - 0.44, 0.5, [P(title, 16, bold=True, color=WHITE)])
        add_text(s, x + 0.22, y0 + 1.48, cw - 0.44, 0.3, [P(en.upper(), 9, bold=True, color=tint(col, 0.55))])
        add_text(s, x + 0.22, y0 + 1.85, cw - 0.44, 1.6, [P(desc, 11.5, color=WHITE, hl=WHITE)])
        if i < n - 1:
            add_text(s, x + cw - 0.02, y0 + h / 2 - 0.2, gap + 0.04, 0.4, [P(CHEVRON, 18, bold=True, color=FAINT, align="c")], anchor="m", wrap=False)
    draw_banner(s, ctx, spec["banner"][0], spec["banner"][1], MX, 6.0, CONTENT_W)
    return s


def k_side(prs, ctx, page, spec):
    s = content_slide(prs, ctx, page, spec)
    col_w = spec.get("col_w", 4.6)
    left_column(s, ctx, spec, col_w)
    rx = MX + col_w + 0.45
    place(s, ctx, spec["rows"], (rx, CONTENT_TOP - 0.1, SLIDE_W - MX - rx, CONTENT_BOTTOM - CONTENT_TOP + 0.1), spec.get("flow", ""))
    return s


def k_wide(prs, ctx, page, spec):
    s = content_slide(prs, ctx, page, spec)
    place(s, ctx, spec["rows"], (MX, 1.9, CONTENT_W, 3.25), spec.get("flow", ""))
    cols = spec["steps"]
    n = len(cols)
    gap = 0.3
    cw = (CONTENT_W - (n - 1) * gap) / n
    y = 5.35
    for i, (title, detail) in enumerate(cols):
        x = MX + i * (cw + gap)
        badge(s, x + BADGE / 2, y + 0.15, i + 1, d=BADGE, size=11.5, fill=ctx.color, ring=False)
        add_text(s, x + BADGE + 0.12, y, cw - BADGE - 0.12, 0.32, [P(title, 13, bold=True, hl=ctx.color)])
        add_text(s, x + BADGE + 0.12, y + 0.34, cw - BADGE - 0.12, 0.85, [P(detail, 10.5, color=MUTED, hl=INK)])
    if spec.get("banner"):
        draw_banner(s, ctx, spec["banner"][0], spec["banner"][1], MX, 6.5, CONTENT_W, 11)
    return s


def k_compare(prs, ctx, page, spec):
    s = content_slide(prs, ctx, page, spec)
    gap = 0.35
    pw = (CONTENT_W - gap) / 2
    bottom = 0.0
    for i, panel in enumerate(spec["panels"]):
        x = MX + i * (pw + gap)
        col = ctx.color if i == 0 else panel.get("color", "2B3140")
        add_rect(s, x, CONTENT_TOP - 0.05, pw, 0.44, fill=col, rounded=0.06)
        add_text(s, x + 0.2, CONTENT_TOP - 0.05, pw - 0.4, 0.44, [P(panel["label"], 13, bold=True, color=WHITE)], anchor="m")
        place(s, ctx, [panel["row"]], (x, CONTENT_TOP + 0.5, pw, 2.45))
        by = CONTENT_TOP + 3.05
        for line in panel["bullets"]:
            lines = count_lines(line, 12, False, pw - 0.35)
            add_text(s, x, by, 0.25, 0.3, [P("•", 13, bold=True, color=col)])
            add_text(s, x + 0.28, by, pw - 0.28, lines * 0.23 + 0.05, [P(line, 12, hl=col)])
            by += lines * 0.23 + 0.1
        bottom = max(bottom, by)
    if spec.get("banner"):
        bh = banner_height(spec["banner"][1], CONTENT_W, 11.5)
        y = min(max(bottom + 0.1, 6.2), 7.05 - bh)
        if y >= bottom:
            draw_banner(s, ctx, spec["banner"][0], spec["banner"][1], MX, y, CONTENT_W, 11.5)
        else:
            print(f"  note: banner skipped on slide {page} (no room)", file=sys.stderr)
    return s


def k_table(prs, ctx, page, spec):
    s = content_slide(prs, ctx, page, spec)
    tw = 7.6
    y = CONTENT_TOP
    for x, w, head in ((MX, 0.5, "#"), (MX + 0.5, 2.9, "الفرع · BRANCH"), (MX + 3.4, 4.2, "المكان في السيستم · LOCATION")):
        add_text(s, x, y, w, 0.25, [P(head, 9, bold=True, color=ctx.color)])
    y += 0.34
    rh = 0.46
    for i, (name, addr) in enumerate(spec["rows_data"], 1):
        add_line(s, MX, y, MX + tw, y, RULE)
        add_text(s, MX, y + 0.11, 0.5, 0.3, [P(str(i), 12, bold=True, color=ctx.color)])
        add_text(s, MX + 0.5, y + 0.11, 2.85, 0.3, [P(name, 12, bold=True)])
        add_text(s, MX + 3.4, y + 0.11, 4.2, 0.3, [P(addr, 11.5, color=MUTED)])
        y += rh
    add_line(s, MX, y, MX + tw, y, RULE)
    rx = MX + tw + 0.5
    rw = SLIDE_W - MX - rx
    place(s, ctx, spec["rows"], (rx, CONTENT_TOP, rw, 2.3))
    by = 4.35
    for b in spec.get("banners", []):
        by += draw_banner(s, ctx, b[0], b[1], rx, by, rw, 11.5) + 0.2
    return s


def k_twice(prs, ctx, page, spec):
    s = content_slide(prs, ctx, page, spec)
    add_text(s, MX, 1.95, 10, 0.4, [P(spec["lead"], 15, color=MUTED, hl=INK)])
    pw, gap = (CONTENT_W - 0.4) / 2, 0.4
    for i, (n, label, main, sub) in enumerate(spec["cards"]):
        x = MX + i * (pw + gap)
        add_rect(s, x, 2.5, pw, 2.25, fill=CARD, line=RULE, line_w=1, rounded=0.1)
        add_rect(s, x, 2.5, 0.09, 2.25, fill=ctx.color)
        add_text(s, x + 0.3, 2.55, 0.9, 1, [P(n, 54, bold=True, color=ctx.color)])
        add_text(s, x + 1.25, 2.72, pw - 1.5, 0.3, [P(label, 11, bold=True, color=ctx.color)])
        add_text(s, x + 1.25, 3.05, pw - 1.5, 0.9, [P(main, 15, bold=True, hl=ctx.color)])
        add_text(s, x + 1.25, 3.95, pw - 1.5, 0.75, [P(sub, 11.5, color=MUTED)])
    steps = spec["flow"]
    n = len(steps)
    g = 0.2
    cw = (CONTENT_W - (n - 1) * g) / n
    for i, t in enumerate(steps):
        x = MX + i * (cw + g)
        strong = i in (0, n - 1)
        add_rect(s, x, 5.0, cw, 1.1, fill=ctx.color if strong else tint(ctx.color, 0.88), rounded=0.08)
        add_text(s, x + 0.15, 5.08, 0.6, 0.4, [P(str(i + 1), 16, bold=True, color=WHITE if strong else ctx.color)])
        add_text(s, x + 0.15, 5.45, cw - 0.3, 0.6, [P(t, 11.5, bold=True, color=WHITE if strong else INK, hl=WHITE if strong else ctx.color)])
        if i < n - 1:
            add_text(s, x + cw - 0.02, 5.33, g + 0.04, 0.4, [P(CHEVRON, 16, bold=True, color=FAINT, align="c")], anchor="m", wrap=False)
    draw_banner(s, ctx, spec["banner"][0], spec["banner"][1], MX, 6.3, CONTENT_W, 11.5)
    return s


def k_checklist(prs, ctx, page, spec):
    ctx.color = ZYDA_BLUE
    s = new_slide(prs, ctx, page, spec.get("notes", ""), bar=ZYDA_BLUE)
    title_block(s, ctx, spec["title"], spec["en"], None)
    items = spec["items"]
    half = math.ceil(len(items) / 2)
    colw = (CONTENT_W - 0.5) / 2
    for idx, (text, sec) in enumerate(items):
        col, row = divmod(idx, half)
        c = SECTIONS[sec]["color"]
        x = MX + col * (colw + 0.5)
        y = CONTENT_TOP + row * 0.98
        add_rect(s, x, y + 0.05, 0.32, 0.32, fill=CARD, line=c, line_w=2, rounded=0.05)
        add_text(s, x + 0.55, y, colw - 0.6, 0.75, [P(text, 14, bold=True, hl=c, after=2),
                                                   P(f"{SECTIONS[sec]['ar']} · SECTION {sec:02d}", 9.5, color=c)])
    return s


KINDS = dict(cover=k_cover, agenda=k_agenda, divider=k_divider, statement=k_statement, journey=k_journey,
             side=k_side, wide=k_wide, compare=k_compare, table=k_table, twice=k_twice, checklist=k_checklist)


# --------------------------------------------------------------------------------------
# The deck.  Marks are (step number, x, y) as fractions of the original screenshot.
# Text: Egyptian Arabic, dashboard UI terms in English exactly as on screen, written **like this**.
# --------------------------------------------------------------------------------------
def deck_spec() -> List[dict]:
    return [
        # ---------------------------------------------------------------- opening
        dict(kind="cover", notes="أهلاً بيكم. السيشن دي هتمشي معاكم على الداشبورد خطوة بخطوة: من الـ Login لحد ما الأوردر يتأكد. كل خطوة ليها screenshot، والأرقام اللي على الصور هي نفس أرقام الخطوات اللي جنبها."),
        dict(kind="agenda", title="هنتكلم في إيه؟", en="What we'll cover",
             notes="خمس أجزاء بنفس ترتيب المكالمة الحقيقية: من أول الـ Login لحد تأكيد الأوردر.",
             items=[
                 (1, "الـ **Login**، أمان الحساب، واختيار البراند الصح"),
                 (2, "رقم الموبايل: عميل جديد ولا مسجّل"),
                 (3, "آخر أوردر، البحث، **Special Instructions**، و **Delivery** ولا **Pickup**"),
                 (4, "مراجعة الأوردر، **Cash** ولا **Online**، الكاش باك والـ **Voucher**"),
                 (5, "الأوردر بيتقبل: الفرع وبيانات العميل ورقم الأوردر"),
             ]),
        dict(kind="journey", title="رحلة الأوردر في 5 خطوات", en="The order journey in 5 steps",
             notes="دي الشغلانة كلها في خمس خطوات. كل جزء من الأجزاء الجاية بيفصّل خطوة منهم.",
             steps=[
                 (1, "Login والبراند", "Login & brand", "ادخل بحسابك واختار البراند اللي العميل طالبه."),
                 (2, "العميل", "Customer", "رقم الموبايل: سجّل الجديد أو افتح بروفايل المسجّل."),
                 (3, "الأوردر", "Order", "راجع آخر أوردر، دوّر على الأصناف، اكتب الملاحظات، وحدد Delivery أو Pickup."),
                 (4, "الدفع", "Checkout", "راجع الأوردر مع العميل، Cash ولا Online، والكاش باك لو موجود."),
                 (5, "التأكيد", "Confirmation", "الأوردر يتقبل ويظهر بالفرع وبيانات العميل ورقم الأوردر."),
             ],
             banner=("tip", "لو تُهت في نص المكالمة، ارجع للسلايد دي.")),

        # ---------------------------------------------------------------- 01 login & brand
        dict(kind="divider", sec=1, items=["تسجيل الدخول", "متشاركش الباسورد", "اختار البراند الصح"],
             notes="أول جزء: الدخول على الداشبورد واختيار البراند اللي العميل بيكلّم عشانه."),
        dict(kind="side", sec=1, title="تسجيل الدخول", en="Log in to the dashboard",
             notes="شاشة الـ Login بتطلب إيميل وباسورد. امشوا على الخانات بالترتيب، وأكّدوا على قاعدة الأمان هنا وفي السلايد الجاية.",
             steps=[
                 ("افتح صفحة الـ **Login**", "هتلاقي عنوان **Login to your store**."),
                 ("اكتب الإيميل بتاعك", "في خانة **Email Address** — الإيميل الخاص بيك إنت بس."),
                 ("اكتب الباسورد", "في خانة **Password**."),
                 ("اضغط **Login**", "هتدخل على الداشبورد."),
             ],
             rows=[[U("login", crop=(0.5, 0.22, 1.0, 0.80), cap="شاشة الـ Login",
                      marks=[(1, 0.60, 0.335), (2, 0.60, 0.477), (3, 0.60, 0.60), (4, 0.60, 0.73)])]],
             banner=("must", "ممنوع تشارك الإيميل أو الباسورد مع أي حد.")),
        dict(kind="statement", sec=1, headline="متشاركش الـ Login بتاعك.", en="Never share your login",
             sub="الإيميل والباسورد ليك إنت بس.",
             chips=[("زمايلك", "Colleagues"), ("المديرين", "Managers"), ("العملاء", "Customers"), ("أي حد تاني", "Anyone else")],
             notes="وقفة هنا. القاعدة بسيطة: بيانات الدخول متتشاركش مع أي حد، مهما كان مين اللي بيطلب ومهما كان السبب."),
        dict(kind="side", sec=1, title="اختار البراند اللي العميل طالبه", en="Select the right brand", flow="h",
             notes="الداشبورد فيه أكتر من براند. قبل ما تسجّل أي حاجة، اتأكد إنك داخل على البراند اللي العميل بيطلب منه، من Switch Store.",
             steps=[
                 ("شوف إنت على أنهي براند", "اسمه ظاهر فوق في القائمة الجانبية — في الصورة **WHAT THE TRUCK**."),
                 ("افتح **Switch Store**", "من اسم البراند اللي فوق."),
                 ("دوّر على البراند", "اكتب اسمه في **Search using store name** أو اختاره من الليستة."),
                 ("اتأكد قبل ما تبدأ", "لازم تشوف **You're now logged in to** واسم البراند الصح."),
             ],
             rows=[[U("store_switcher", cap="البراند الحالي", marks=[(1, 0.64, 0.30)]),
                    U("switch_store", crop=(0.31, 0.02, 0.69, 1.0), cap="Switch Store",
                      marks=[(2, 0.45, 0.14), (4, 0.63, 0.27), (3, 0.63, 0.395)])]],
             banner=("warn", "اتأكد من اسم البراند قبل ما تسجّل أي حاجة للعميل.")),

        # ---------------------------------------------------------------- 02 customer data
        dict(kind="divider", sec=2, items=["رقم الموبايل", "عميل جديد", "عميل مسجّل"],
             notes="تاني جزء: جوه البراند اللي اخترناه بنسجّل بيانات العميل."),
        dict(kind="side", sec=2, title="الخطوة 1 · اكتب رقم الموبايل", en="Enter the phone number", flow="v",
             notes="رقم الموبايل هو اللي السيستم بيعرف بيه العميل. اكتبه بالراحة، وبعدين شوف السيستم قال إيه: يا العميل يظهر، يا يقول إنه عميل جديد.",
             steps=[
                 ("افتح **Select Customer**", "هيطلب منك **Phone Number** بتاع العميل."),
                 ("اكتب رقم الموبايل", "الخانة بتبدأ بكود مصر **\u200e+20**. كمّل الرقم وراجعه رقم رقم."),
                 ("شوف السيستم قال إيه", "العميل المسجّل بيظهر على طول. الرقم الجديد بيطلع **This looks like a new customer**."),
             ],
             rows=[[U("phone_entry", marks=[(1, 0.60, 0.12), (2, 0.70, 0.67)])],
                   [U("customer_new", cap="بعد رقم مش متسجّل", marks=[(3, 0.88, 0.68)])]],
             banner=("tip", "اقرأ الرقم للعميل تاني قبل ما تكمل.")),
        dict(kind="side", sec=2, title="عميل جديد: سجّله", en="New customer: register them", flow="h",
             notes="الرقم الجديد محتاج نسجّل العميل. المطلوب هنا الاسم بس. زرار Save بيفضل رمادي لحد ما الاسم يبقى صح.",
             steps=[
                 ("السيستم بيقول إنه عميل جديد", "**This looks like a new customer** — مفيش بيانات للرقم ده."),
                 ("اضغط **Add Customer**", None),
                 ("اكتب الاسم في **Full Name**", "حروف بس. لو فيه أرقام هيطلع **Name should contain letters** و **Save** هيفضل رمادي."),
                 ("اضغط **Save**", "العميل اتسجّل وتقدر تكمّل الأوردر."),
             ],
             rows=[[U("customer_new", marks=[(1, 0.88, 0.68), (2, 0.10, 0.83)]),
                    U("customer_details", marks=[(3, 0.75, 0.55)])],
                   [U("customer_save", cap="Save بيشتغل لما الاسم يبقى صح", marks=[(4, 0.07, 0.5)])]]),
        dict(kind="compare", sec=2, title="عميل جديد ولا مسجّل؟", en="New vs existing customer",
             notes="الفرق كله في الأول: العميل الجديد بنسجّله، والمسجّل السيستم بيعرفه وبنفتح بروفايله، وعناوينه المحفوظة بتبقى جاهزة في الـ Checkout.",
             panels=[
                 dict(label="عميل جديد · New", row=[U("customer_new")], bullets=[
                     "السيستم بيقول **This looks like a new customer**.",
                     "اضغط **Add Customer**، اكتب **Full Name**، و **Save**.",
                     "مفيش عناوين محفوظة — هتضيف العنوان في الـ **Checkout**.",
                 ]),
                 dict(label="عميل مسجّل · Existing", row=[
                     U("customer_existing", cap="عميل مسجّل"),
                     U("saved_addresses", crop=(0, 0.56, 1, 1), cap="العناوين المحفوظة")], bullets=[
                     "الرقم بيتعرف على طول — من غير تسجيل، افتح البروفايل وكمّل.",
                     "اتأكد من الاسم مع العميل.",
                     "عناوينه بتظهر في **Select Address**: اختار واحد أو ضيف جديد.",
                 ]),
             ]),

        # ---------------------------------------------------------------- 03 order processing
        dict(kind="divider", sec=3, items=["آخر أوردر", "البحث", "Special Instructions", "Delivery ولا Pickup"],
             notes="تالت جزء: تنفيذ الأوردر نفسه."),
        dict(kind="side", sec=3, title="راجع آخر أوردر للعميل", en="Check the last order",
             notes="قبل ما نبدأ الأوردر الجديد: آخر أوردر للعميل بيظهر فوق الـ Categories. اسأل العميل لو كان فيه أي مشكلة فيه قبل ما تكمّل.",
             steps=[
                 ("شوف آخر أوردر", "بيظهر فوق الـ **Categories** في المنيو."),
                 ("اسأل العميل عنه", "كان فيه أي مشكلة في آخر أوردر؟"),
                 ("وبعدين كمّل", "ابدأ الأوردر الجديد."),
             ],
             rows=[[U("last_order")]],
             banner=("tip", "السؤال ده قبل أي أوردر جديد لعميل مسجّل.")),
        dict(kind="side", sec=3, title="دوّر على الصنف بالـ Search", en="Find items fast: search", col_w=4.7,
             notes="الـ Search أسرع طريقة. مش لازم تكتب الاسم كله: أول كام حرف والسيستم يطلع الصنف. تحت كل اسم فيه وصف تفصيلي.",
             steps=[
                 ("اضغط على **Search**", "فوق المنيو."),
                 ("اكتب أول كام حرف", "مثلاً لما تكتب **RAN** السيستم بيطلع **Ranch** و **Creamy Chicken Ranch** على طول."),
                 ("اقرأ الوصف تحت الاسم", "كل صنف تحته وصف تفصيلي — استخدمه لو العميل بيسأل."),
                 ("اتأكد إنه متاح واختاره", "لو عليه **Sold out** يبقى مش متاح. لو متاح اضغط عليه."),
             ],
             rows=[[U("search_first_letters", cap="أول كام حرف: RAN",
                      marks=[(1, 0.75, 0.035), (2, 0.55, 0.17), (2, 0.62, 0.395), (3, 0.70, 0.47)]),
                    U("search_results", cap="نتايج البحث عن ice", boxes=[(0.03, 0.195, 0.74, 0.228)],
                      marks=[(3, 0.55, 0.248), (4, 0.88, 0.14)])]],
             banner=("tip", "مش لازم تكتب الاسم كله — أول كام حرف كفاية.")),
        dict(kind="side", sec=3, title="افتح الصنف واختار الـ Options", en="Open an item and choose its options", col_w=4.9,
             notes="افتح الصنف، اختار الحجم، جاوب على أي جروب Required، حدد الكمية، وبعدين Add Item. السعر على الزرار بيمشي مع اختياراتك.",
             steps=[
                 ("افتح الصنف", "الصورة والاسم والوصف فوق."),
                 ("اختار الـ **Size**", "كل حجم ليه سعره: في المثال ده Regular بـ EGP 160 و All Day بـ EGP 280."),
                 ("جاوب على الـ **Required**", "أي جروب مكتوب عليه **Required** لازم تختار منه."),
                 ("حدد الكمية", "بـ − و + تحت."),
                 ("اضغط **Add Item**", "السعر على الزرار بيتغير حسب اختياراتك."),
             ],
             inset=[U("basket_items", cap="بعد Add Item: الصنف في Your Items باختياره")],
             rows=[[U("item_options", marks=[(1, 0.75, 0.47), (2, 0.60, 0.65), (3, 0.65, 0.795), (4, 0.12, 0.905), (5, 0.62, 0.95)])]]),
        dict(kind="side", sec=3, title="طلبات العميل: Special Instructions", en="Comments & exclusions", flow="",
             notes="بعض الأصناف تحتها خانة Special Instructions. فيها بنكتب طلبات العميل: حاجة عايز يشيلها زي الصوص أو الخس أو الخضار، أو أي تفضيل. الإضافات اللي بسعر بتتختار من جروب Extra.",
             steps=[
                 ("افتح الصنف", "بعض الأصناف بس تحتها خانة ملاحظات."),
                 ("اكتب في **SPECIAL INSTRUCTIONS**", "طلب العميل أو الحاجة اللي عايز يشيلها: من غير صوص، من غير خس، من غير خضار."),
                 ("الإضافات بسعر من **Extra**", "اختارها من الـ options — بتتسجل على الأوردر تحت **Extra** بسعرها."),
                 ("اضغط **Add Item**", "الصنف بيتضاف بالملاحظة بتاعته."),
             ],
             rows=[[U("special_instructions", boxes=[(0.04, 0.725, 0.93, 0.81)], marks=[(1, 0.35, 0.61), (2, 0.88, 0.785)]),
                    U("order_item_options", crop=(0, 0.26, 1, 0.70), cap="الإضافات بسعر على الأوردر: Extra", marks=[(3, 0.50, 0.64)])]],
             banner=("note", "خانة **Special Instructions** موجودة تحت أصناف معينة بس — مش كل الأصناف.")),
        dict(kind="side", sec=3, title="الأصناف الـ Sold out والمش متاحة", en="Sold out or not available", col_w=5.0,
             notes="الأصناف المش متاحة بتبان في المنيو بس بعلامة واضحة: رمادي ومكتوب Sold out. الأصناف اللي محتاجة ميعاد عليها Schedule for. قول للعميل على طول واقترح بديل.",
             steps=[
                 ("صنف متاح", "ألوانه عادية وعلى الصورة زرار **+**."),
                 ("**Sold out**", "الاسم والسعر والصورة رمادي ومكتوب **Sold out**. مفيش زرار **+**."),
                 ("بميعاد بس", "علامة زي **Schedule for Sep 28** معناها إن الصنف مش متاح **ASAP** — متاح في اليوم ده."),
             ],
             inset=[U("sold_out_item", cap="نفس علامة Sold out على صنف تاني", marks=[(2, 0.30, 0.78)])],
             rows=[[U("sold_out_list", marks=[(1, 0.55, 0.14), (2, 0.25, 0.385), (3, 0.42, 0.572), (2, 0.25, 0.73)])]],
             banner=("warn", "قول للعميل على طول واقترح بديل من نفس الكاتيجوري.")),
        dict(kind="compare", sec=3, title="Delivery ولا Pickup؟", en="Delivery or pickup?",
             notes="بعد ما نختار الأصناف لازم نعرف الأوردر Delivery ولا Pickup. الـ Delivery اللوكيشن بيتاخد فيه مرتين، والـ Pickup بنختار فيه الفرع والميعاد.",
             panels=[
                 dict(label="Delivery · توصيل", row=[U("map_empty", crop=(0, 0, 1, 0.32), cap="Order Mode: Delivery", marks=[(1, 0.12, 0.095)])], bullets=[
                     "اختار **Delivery** في **Order Mode**.",
                     "اللوكيشن بيتاخد **مرتين**: في الأول وفي الـ **Checkout**.",
                     "التفاصيل في السلايدات الجاية.",
                 ]),
                 dict(label="Pickup · استلام من الفرع", row=[U("pickup_time", cap="Pickup", marks=[(1, 0.45, 0.14)])], bullets=[
                     "اختار **Pickup**.",
                     "اختار الفرع من ليستة الفروع.",
                     "حدد ميعاد الاستلام: **ASAP** أو **Schedule slot**.",
                 ]),
             ]),
        dict(kind="twice", sec=3, title="الـ Delivery: اللوكيشن بيتاخد مرتين", en="The location is taken twice",
             notes="دي القاعدة اللي ناس كتير بتنساها: اللوكيشن بيتكتب مرتين. مرة على خريطة Order Mode، ومرة تانية في الـ Checkout من Add new address مع تفاصيل العنوان. امشوا على الخمس خطوات اللي تحت.",
             lead="كل أوردر **Delivery** محتاج لوكيشن العميل يتحط مرتين.",
             cards=[
                 ("1", "في الأول · AT THE START", "حط لوكيشن العميل على خريطة **Order Mode** واضغط **Save**.",
                  "بيحدد منطقة التوصيل والرسوم والوقت — مثلاً El Shorouk – 5th District برسوم EGP 30.00 ووقت من 35 لـ 50 دقيقة."),
                 ("2", "في الـ Checkout · AT CHECKOUT", "اضغط **Add new address**، حط اللوكيشن تاني، واكتب تفاصيل العنوان واحفظ.",
                  "كده العنوان اتسجّل على العميل."),
             ],
             flow=["حط اللوكيشن الأول", "اختار صنف وهمي", "روح للـ **Checkout**", "اضغط **Add new address**", "حط اللوكيشن تاني وسجّل العنوان"],
             banner=("must", "خد اللوكيشن مرتين: مرة في الأول، ومرة من **Add new address** في الـ **Checkout**.")),
        dict(kind="side", sec=3, title="الـ Delivery · خطوة 1 · حط اللوكيشن الأول", en="Enter the initial location", flow="h",
             notes="ده أول مرة بناخد فيها اللوكيشن. حط لوكيشن العميل على خريطة Order Mode واحفظ. لو السيستم قال We don't deliver to this address يبقى الـ pin برّه منطقة التوصيل.",
             steps=[
                 ("اختار **Delivery**", "في **Order Mode**، أول تاب."),
                 ("دوّر على اللوكيشن", "اكتب العنوان أو الصق الإحداثيات في **Enter Location**."),
                 ("اختار الاقتراح", "الـ pin بيروح عليه. قبلها الزرار مكتوب عليه **Move The Pin**."),
                 ("اضغط **Save**", "كده اللوكيشن اتاخد (أول مرة)."),
             ],
             inset=[U("not_deliverable")],
             rows=[[U("map_search", marks=[(1, 0.10, 0.095), (2, 0.92, 0.18), (3, 0.92, 0.275)]),
                    U("map_save", marks=[(4, 0.12, 0.95)])]]),
        dict(kind="side", sec=3, title="الـ Delivery · خطوة 2 · اختار صنف وهمي", en="Pick a fake item", flow="h",
             notes="الـ Checkout مش بيفتح غير لما يبقى فيه حاجة في الباسكت. فبنختار أي صنف كبديل مؤقت، بس عشان نوصل للـ Checkout اللي بنسجّل منه العنوان.",
             steps=[
                 ("اختار أي صنف", "بالـ Search أو من المنيو — هو بس بديل مؤقت."),
                 ("افتحه وضيفه", "اضغط على الصنف وبعدين **Add Item**."),
                 ("اضغط **Checkout**", "الزرار بيظهر في الشريط اللي تحت لما الباسكت يبقى فيه صنف."),
             ],
             inset=[U("checkout_bar")],
             rows=[[U("search_results", cap="اختار أي صنف", marks=[(1, 0.80, 0.04), (2, 0.88, 0.14)]),
                    U("basket_checkout", marks=[(3, 0.62, 0.93)])]],
             banner=("tip", "الصنف الوهمي بس عشان يفتح الـ **Checkout** — العنوان بيتسجّل من الـ **Checkout**.")),
        dict(kind="side", sec=3, title="الـ Delivery · خطوة 3 · Add new address", en="Checkout: add new address", flow="h",
             notes="في الـ Checkout، جزء Deliver to مكتوب فيه Please add an address بالأحمر. اضغط عليه يفتح Select Address، وبعدين Add new address.",
             steps=[
                 ("بص على **Deliver to**", "مكتوب **Please add an address** بالأحمر: لسه مفيش عنوان على الأوردر."),
                 ("اضغط على **Deliver to**", "هتفتح **Select Address**، ولو فيه عناوين محفوظة هتظهر هنا."),
                 ("اضغط **Add new address**", "دي بتفتح شاشة اللوكيشن تاني مرة."),
             ],
             rows=[[U("checkout_add_address", marks=[(1, 0.70, 0.86), (2, 0.93, 0.70)]),
                    U("saved_addresses", cap="Select Address", marks=[(3, 0.10, 0.70)])]]),
        dict(kind="side", sec=3, title="الـ Delivery · خطوة 4 · سجّل العنوان", en="Location again + register the address", flow="h", col_w=4.8,
             notes="دي تاني مرة بناخد فيها اللوكيشن. حط الـ pin تاني، وبعدين املا فورم العنوان واحفظ، كده العنوان اتسجّل على العميل.",
             steps=[
                 ("حط اللوكيشن تاني", "الخريطة بتفتح تاني مرة. دوّر على نفس اللوكيشن وحط الـ pin عليه."),
                 ("اضغط **Save**", "فورم العنوان هيفتح."),
                 ("اختار النوع", "**House** أو **Apartment** أو **Office**."),
                 ("راجع **City** و **Area**", "متعبّية لوحدها — اتأكد إنها مظبوطة."),
                 ("املا التفاصيل", "**Street** و **Building** و **Floor** و **Apartment Number**. الـ **Notes** و **Save Address As** اختياري."),
                 ("اضغط **Add new address**", "العنوان اتسجّل على العميل."),
             ],
             rows=[[U("map_second_location", marks=[(1, 0.72, 0.255), (2, 0.12, 0.935)]),
                    U("address_form", marks=[(3, 0.36, 0.255), (4, 0.88, 0.352), (5, 0.92, 0.55), (6, 0.10, 0.955)])]],
             banner=("tip", "الـ pin في مكان غلط؟ اضغط **Edit Location** فوق الفورم.")),
        dict(kind="side", sec=3, title="الـ Pickup · اختار الفرع", en="Pickup: choose the branch", flow="h",
             notes="للـ Pickup، غيّر Order Mode لـ Pickup. ليستة Select Branch فيها كل الفروع، الأقرب الأول، بالمسافة والعنوان. اضغط على السهم بتاع الفرع الصح.",
             steps=[
                 ("اختار **Pickup**", "في **Order Mode** جنب **Delivery**."),
                 ("شوف ليستة الفروع", "**Select Branch** فيها كل الفروع، والأقرب الأول — مكتوب **Sorted by distance**."),
                 ("اضغط على سهم الفرع", "كل سطر فيه اسم الفرع والمسافة والعنوان."),
                 ("راجع كارت الفرع", "الفرع المختار بيظهر كـ **Pickup Branch** بالـ pin بتاعه."),
             ],
             rows=[[U("pickup_branches", marks=[(1, 0.60, 0.09), (2, 0.64, 0.14), (3, 0.82, 0.31)]),
                    U("pickup_branch_card", marks=[(4, 0.65, 0.86)])]],
             banner=("tip", "ليستة الفروع كاملة في السلايد الجاية.")),
        dict(kind="table", sec=3, title="الفروع فين؟", en="Where are the branches?",
             notes="دي الليستة اللي السيستم بيعرضها. المسافات بتتغير حسب لوكيشن العميل، فاقروها من الشاشة. الليستة بتعمل scroll، فممكن يكون فيه فروع تانية تحت Golf Central.",
             rows_data=[
                 ("Street Side Mall - El Shorouk", "El Shorouk – Street Side Mall"),
                 ("L4, The Strip Mall, Madinaty", "The Strip, Madinaty"),
                 ("The Drive 2", "The Drive 2"),
                 ("The Yard Mall Rehab", "The Yard Mall Gate 6 Rehab"),
                 ("The Mind Space", "Abdullah ibn Salamah, st, New Cairo 1, Cairo"),
                 ("Sway Mall", "Sway mall"),
                 ("Promenade Mall New Cairo", "Promenade Mall, Narges, 79 Street"),
                 ("District 5", "D5 Mall at District 5 Compound, Cairo Governorate"),
                 ("Golf Central - Palm Hills", "Palm Hills"),
             ],
             rows=[[U("pickup_branch_card", cap="كل فرع ليه pin على الخريطة")]],
             banners=[("note", "الليستة بتعمل scroll — ممكن يكون فيه فروع تانية تحت Golf Central."),
                      ("tip", "المسافة بتتغير حسب لوكيشن العميل. اقراها من الشاشة.")]),
        dict(kind="side", sec=3, title="الـ Pickup · ميعاد الاستلام", en="Pickup date and time",
             notes="ميعاد الـ Pickup الافتراضي ASAP مع وقت تقريبي (Ready in 15 minutes). لو عايز وقت تاني، Schedule slot واختار اليوم والساعة. الميعاد ده بيتسجّل على الأوردر.",
             steps=[
                 ("**Pickup** مختار", "التوجل عليه **Pickup**."),
                 ("راجع الفرع", "**Picking up from** فيه اسم الفرع. **Change** لو عايز تغيّره."),
                 ("الميعاد الافتراضي", "**Ready in 15 minutes - ASAP**."),
                 ("عايز ميعاد تاني؟ **Schedule slot**", "اختار يوم وساعة الاستلام."),
                 ("الميعاد بيتسجّل على الأوردر", "الفرع ويوم وساعة الاستلام بيظهروا على كارت الأوردر."),
             ],
             rows=[[U("pickup_time", marks=[(1, 0.45, 0.14), (2, 0.60, 0.42), (3, 0.50, 0.72), (4, 0.68, 0.73)])],
                   [U("schedule_slot_picker"), U("pickup_order_recorded")]]),

        # ---------------------------------------------------------------- 04 checkout & payment
        dict(kind="divider", sec=4, items=["راجع مع العميل", "Cash ولا Online", "الكاش باك", "الـ Voucher"],
             notes="رابع جزء: مراجعة الأوردر مع العميل والدفع."),
        dict(kind="side", sec=4, title="راجع الأوردر مع العميل", en="Review the order with the customer", flow="h",
             notes="قبل الدفع: راجع مع العميل تفاصيل الأوردر والـ Total وبياناته.",
             steps=[
                 ("راجع الأصناف", "كل صنف بالـ options والملاحظات بتاعته."),
                 ("راجع الـ **Total**", "فيه **Subtotal** و **Service Fee**، و **Delivery Services** في الـ Delivery بس."),
                 ("راجع بيانات العميل", "الاسم والموبايل والعنوان، أو الفرع لو **Pickup**."),
                 ("خد موافقة العميل", "قبل ما تكمّل للدفع."),
             ],
             rows=[[U("basket_checkout", cap="Your Items", marks=[(1, 0.08, 0.40)]),
                    U("checkout_full", boxes=[(0.01, 0.67, 0.98, 0.90)], marks=[(2, 0.75, 0.875), (3, 0.70, 0.34)])]]),
        dict(kind="compare", sec=4, title="اسأل العميل: Cash ولا Online؟", en="Payment: cash or online",
             notes="اسأل العميل الأول هيدفع إزاي. Cash: الزرار Place Order. Online هو Credit Card: الزرار بيتغير لـ Send Cart Link، وفي براندات معينة بيظهر سطر Cashback.",
             panels=[
                 dict(label="Cash · كاش", row=[U("payment_place_order", boxes=[(0.0, 0.22, 1.0, 0.45)])], bullets=[
                     "اختار **Cash** في **Pay with**.",
                     "العميل هيدفع كاش.",
                     "الزرار بيبقى **Place Order**.",
                 ]),
                 dict(label="Online · Credit Card", row=[U("cashback_checkout", boxes=[(0.02, 0.53, 0.95, 0.65)])], bullets=[
                     "اختار **Credit Card** — ده الدفع الـ **Online**.",
                     "الزرار بيتغير لـ **Send Cart Link** — العميل بيدفع أونلاين من اللينك.",
                     "في براندات معينة بيظهر **Cashback** — التفاصيل في السلايد الجاية.",
                 ]),
             ],
             banner=("tip", "اسأل العميل هيدفع إزاي الأول، وبعدين اختار.")),
        dict(kind="side", sec=4, title="الكاش باك: للدفع الـ Online بس", en="Cashback: online payment only", col_w=4.7,
             notes="الكاش باك موجود في تلات براندات بس: Maine و Vinnys Pizza و Chickin Worx، وبشرط إن العميل يدفع Online. لما تختار Credit Card هيظهر سطر Cashback بالمبلغ.",
             steps=[
                 ("في 3 براندات بس", "**Maine** و **Vinnys Pizza** و **Chickin Worx**."),
                 ("لازم الدفع يكون Online", "اختار **Credit Card** — مع الـ **Cash** مفيش كاش باك."),
                 ("هتشوف علامة الكاش باك", "سطر **Cashback** بالمبلغ تحت **Pay with** — في المثال ده EGP 57.60."),
                 ("عرّف العميل", "لو في براند من التلاتة وعايز يدفع كاش، قوله إن الكاش باك للـ **Online** بس."),
             ],
             rows=[[U("switch_store", crop=(0.31, 0.44, 0.69, 0.99), cap="البراندات اللي فيها كاش باك",
                      boxes=[(0.325, 0.455, 0.665, 0.525), (0.325, 0.565, 0.665, 0.635), (0.325, 0.89, 0.665, 0.96)],
                      marks=[(1, 0.50, 0.49), (1, 0.50, 0.60), (1, 0.50, 0.925)]),
                    U("cashback_checkout", boxes=[(0.02, 0.71, 0.95, 0.81)],
                      marks=[(2, 0.60, 0.59), (3, 0.55, 0.76)])]],
             banner=("must", "الكاش باك في **Maine** و **Vinnys Pizza** و **Chickin Worx** بس — وبالدفع الـ **Online** بس.")),
        dict(kind="side", sec=4, title="ضيف الـ Voucher", en="Add a discount voucher",
             notes="الـ Voucher مكانه في خانة Save on this order في الـ Checkout، بين ميعاد التوصيل والـ Order Summary. اكتب الكود واضغط Apply.",
             steps=[
                 ("روح للـ **Checkout**", "الصفحة كلها في الصورة."),
                 ("دوّر على **Save on this order**", "بين **Arrives in** والـ **Order Summary**."),
                 ("اكتب الكود في **Voucher**", "زي ما العميل قاله بالظبط."),
                 ("اضغط **Apply**", "الخصم بيظهر في الـ **Order Summary** وعلى الأوردر كـ **Coupon Discount**."),
             ],
             inset=[U("checkout_bar", cap="الشريط ده بيقولك محتاج تزوّد قد إيه عشان الخصم")],
             rows=[[U("checkout_full", boxes=[(0.01, 0.52, 0.98, 0.665)],
                      marks=[(1, 0.93, 0.045), (2, 0.55, 0.552), (3, 0.45, 0.60), (4, 0.92, 0.552)])]],
             banner=("tip", "ضيف الـ Voucher قبل ما تضغط **Place Order**.")),

        # ---------------------------------------------------------------- 05 confirmation
        dict(kind="divider", sec=5, items=["خلّص الـ Checkout", "الأوردر بيوصل الداشبورد", "الفرع وبيانات العميل ورقم الأوردر"],
             notes="آخر جزء: نخلّص الـ Checkout والأوردر يتقبل، ونراجع تفاصيله."),
        dict(kind="side", sec=5, title="خلّص الـ Checkout", en="Complete the checkout", flow="h",
             notes="راجع الـ Order Summary. Cash: Place Order. Online: Send Cart Link والعميل يكمّل الدفع من اللينك. بعدها الأوردر بيتقبل ويظهر على الداشبورد.",
             steps=[
                 ("راجع الـ **Order Summary**", "الـ **Delivery** بيزوّد سطر **Delivery Services**، الـ **Pickup** لأ."),
                 ("Cash: اضغط **Place Order**", None),
                 ("Online: اضغط **Send Cart Link**", "العميل بيكمّل الدفع من اللينك."),
                 ("الأوردر بيتقبل", "ويظهر على الداشبورد — السلايدات الجاية."),
             ],
             inset=[U("summary_delivery", marks=[(1, 0.62, 0.52)])],
             rows=[[U("payment_place_order", cap="Cash", marks=[(2, 0.10, 0.86)]),
                    U("cashback_checkout", cap="Online", marks=[(3, 0.10, 0.91)])]]),
        dict(kind="wide", sec=5, title="الأوردر بيوصل الداشبورد", en="Orders drop into the dashboard",
             notes="كل أوردر بيوصل ككارت. الصف اللي فوق فيه الملخص: الوقت، رقم الأوردر، العميل، Delivery ولا Pickup، الدفع والمبلغ، الميعاد، والحالة.",
             rows=[[U("orders_incoming", crop=(0, 0, 1, 0.64), cap="الأوردر وهو بيوصل",
                      marks=[(1, 0.022, 0.064), (2, 0.955, 0.064), (3, 0.36, 0.30), (4, 0.74, 0.205)])]],
             steps=[
                 ("الأوردر وصل", "كل أوردر كارت لوحده. فوق: الوقت، رقم الأوردر، العميل، النوع، الدفع والمبلغ، الميعاد."),
                 ("الحالة", "البادج اللي فوق في الصورة على اليمين — مثلاً **Accepted**."),
                 ("الأصناف والحساب", "شمال الصورة: كل صنف بالـ options، والـ Subtotal والكوبون والـ VAT والـ Total."),
                 ("العميل والعنوان", "يمين الصورة: الاسم والموبايل والعنوان والملاحظة والخريطة."),
             ],
             banner=("tip", "احفظ شكل الكارت ده — هتراجع عليه كل أوردر بتعمله.")),
        dict(kind="side", sec=5, title="راجع الأوردر بعد ما يتقبل", en="Confirm the accepted order", col_w=3.9,
             notes="الأوردر المقبول بيعرض اسم الفرع وبيانات العميل ورقم الأوردر. راجع التلاتة.",
             steps=[
                 ("اسم الفرع", "مكتوب **Order Accepted by** واسم الفرع — في المثال ده Sway mall."),
                 ("بيانات العميل", "الاسم والموبايل والعنوان."),
                 ("رقم الأوردر", "مكتوب تحت اسم العميل."),
                 ("الأصناف والحساب", "كل صنف باختياراته والـ **Total**."),
             ],
             rows=[[U("order_full", marks=[(1, 0.80, 0.88), (2, 0.80, 0.13), (3, 0.66, 0.065), (4, 0.36, 0.40)])]],
             banner=("must", "اتأكد إن الفرع وبيانات العميل ورقم الأوردر كلهم صح.")),

        # ---------------------------------------------------------------- close
        dict(kind="checklist", title="قبل ما تقفل الأوردر", en="Before you close the order",
             notes="نختم بالـ checklist دي. هي الديك كله في صفحة واحدة.",
             items=[
                 ("دخلت بحسابي — ومشاركتش الباسورد.", 1),
                 ("البراند صح: مكتوب **You're now logged in to** واسم البراند.", 1),
                 ("رقم الموبايل: الجديد اتسجّل، والمسجّل اتفتح بروفايله.", 2),
                 ("راجعت آخر أوردر وسألت العميل لو كان فيه مشكلة.", 3),
                 ("الأصناف صح، والملاحظات في **Special Instructions**.", 3),
                 ("**Delivery**: اللوكيشن اتاخد مرتين. **Pickup**: الفرع والميعاد صح.", 3),
                 ("راجعت الأوردر والـ **Total** وبيانات العميل — والـ **Voucher** لو فيه.", 4),
                 ("**Cash** أو **Online** — والكاش باك في Maine و Vinnys Pizza و Chickin Worx بس.", 4),
                 ("الأوردر اتقبل: الفرع وبيانات العميل ورقم الأوردر صح.", 5),
             ]),
    ]


# --------------------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------------------
def build(ctx: Ctx, out_path: str) -> None:
    prs = Presentation()
    prs.slide_width = I(SLIDE_W)
    prs.slide_height = I(SLIDE_H)
    prs.core_properties.title = "Dashboard Training - step by step"
    prs.core_properties.subject = "Customer-service agent training: ordering dashboard"
    specs = deck_spec()
    for i, sp in enumerate(specs, 1):
        sec = sp.get("sec")
        if sec and sec not in ctx.section_first_page:
            ctx.section_first_page[sec] = i
    for i, sp in enumerate(specs, 1):
        KINDS[sp["kind"]](prs, ctx, i, sp)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)


def list_slots() -> None:
    print(f"{'file name (any of .png/.jpg/.jpeg)':38} what it should show")
    for k, f, size, cap, show in SHOT_TABLE:
        print(f"{f:38} {show}")


def main(argv=None) -> int:
    global RTL, ARROW, CHEVRON
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--screenshots", default=os.path.join(here, "screenshots"), help="folder with the screenshot files")
    ap.add_argument("--out", default=os.path.join(here, "output", "Zyda_Dashboard_Training.pptx"), help="output .pptx")
    ap.add_argument("--header", default=RUNNING_HEADER, help="running header text (default: %(default)s; '' to hide)")
    ap.add_argument("--logo", default=None, help="optional logo image; otherwise the wordmark is drawn as vector shapes")
    ap.add_argument("--placeholders-only", action="store_true", help="ignore any screenshots and draw every slot as a placeholder")
    ap.add_argument("--ltr", action="store_true", help="left-to-right layout (the text stays Arabic + English)")
    ap.add_argument("--list", action="store_true", help="list every screenshot slot and exit")
    args = ap.parse_args(argv)
    if args.list:
        list_slots()
        return 0
    if args.ltr:
        RTL, ARROW, CHEVRON = False, "→", "›"
    ctx = Ctx(args.screenshots, args.placeholders_only, args.header, args.logo)
    build(ctx, args.out)
    print(f"Wrote {args.out}")
    print(f"  screenshots placed : {len(ctx.found)}")
    print(f"  placeholders drawn : {len(ctx.missing)}")
    for k in ctx.missing:
        print(f"    - {SHOTS[k]['file']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
